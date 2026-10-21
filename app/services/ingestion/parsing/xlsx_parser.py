"""XlsxDocumentParser: Excel workbooks as sheets of titled tables (openpyxl).

Each sheet becomes a heading; inside it, runs of non-empty rows separated by
blank rows are blocks. A block whose rows have several filled cells is a
table (its first row the header, kept verbatim), anything else is text
(report titles like "ĐÁNH GIÁ KẾT QUẢ NHẬN DIỆN — Từ 09/10 đến 21/10",
section labels like "Quầy line"). Cached formula results are read
(``data_only``), so a summary sheet shows its numbers, not "=A1/B1".

Large tables are split into row groups by the chunker, each repeating the
header row, so every chunk stays readable on its own.
"""

from __future__ import annotations

import asyncio
import io
from datetime import date, datetime, time
from typing import Any

from app.domain.models.ingestion.documents import SourceDocument
from app.domain.models.ingestion.parsed_document import (
    HeadingNode,
    ParagraphNode,
    ParsedDocument,
    ParsedDocumentMetadata,
    SourceAnchor,
    TableNode,
)
from app.services.ingestion.parsing.base import DocumentParseError, ParserUnavailableError

XLSX_PARSER_NAME = "openpyxl"
MAX_COLUMNS = 40
MAX_ROWS_PER_SHEET = 5000


def openpyxl_version() -> str:
    try:
        import openpyxl

        return str(openpyxl.__version__)
    except ImportError:  # pragma: no cover
        return "unknown"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.strftime("%Y-%m-%d %H:%M:%S")
            if value.time() != time(0)
            else value.date().isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6g}"
    text = str(value).replace("\r", " ").replace("\n", " / ")
    return " ".join(text.split()).replace("|", "\\|")


def _blocks(rows: list[list[str]]) -> list[list[list[str]]]:
    """Runs of non-empty rows, separated by blank rows."""
    blocks: list[list[list[str]]] = []
    current: list[list[str]] = []
    for row in rows:
        if any(row):
            current.append(row)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def _filled(row: list[str]) -> int:
    return sum(1 for cell in row if cell)


class XlsxDocumentParser:
    async def parse(self, source: SourceDocument, content: bytes) -> ParsedDocument:
        return await asyncio.to_thread(self._parse_sync, source, content)

    def _parse_sync(self, source: SourceDocument, content: bytes) -> ParsedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ParserUnavailableError("openpyxl is not installed") from exc
        try:
            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError(f"openpyxl failed on {source.filename}: {exc}") from exc

        nodes: list[HeadingNode | ParagraphNode | TableNode] = []
        warnings: list[str] = []

        def add(node_cls: Any, **kw: Any) -> None:
            nodes.append(node_cls(node_id=f"n{len(nodes):04d}", order=len(nodes), **kw))

        for sheet_index, ws in enumerate(wb.worksheets):
            rows: list[list[str]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS_PER_SHEET:
                    warnings.append(f"sheet '{ws.title}' truncated at {MAX_ROWS_PER_SHEET} rows")
                    break
                rows.append([_cell_text(v) for v in row[:MAX_COLUMNS]])
            if not any(any(r) for r in rows):
                continue
            sheet = str(ws.title).strip() or f"Sheet {sheet_index + 1}"
            anchor = SourceAnchor(block_index=sheet_index)
            add(HeadingNode, text=sheet, heading_path=(), anchor=anchor, level=1)
            path = (sheet,)
            for block in _blocks(rows):
                width = max(len(r) for r in block)
                used = [c for c in range(width) if any(c < len(r) and r[c] for r in block)]
                trimmed = [[r[c] if c < len(r) else "" for c in used] for r in block]
                # Single-cell rows on top ("Quầy line" above its metrics table)
                # are labels, not the table header.
                labels: list[list[str]] = []
                while len(trimmed) > 2 and _filled(trimmed[0]) <= 1:
                    labels.append(trimmed.pop(0))
                if labels:
                    label_text = "\n".join(" — ".join(c for c in r if c) for r in labels)
                    add(ParagraphNode, text=label_text, heading_path=path, anchor=anchor)
                if len(trimmed) >= 2 and _filled(trimmed[0]) >= 2:
                    header = [h or f"Cột {i + 1}" for i, h in enumerate(trimmed[0])]
                    lines = [
                        "| " + " | ".join(header) + " |",
                        "|" + "---|" * len(header),
                        *("| " + " | ".join(r) + " |" for r in trimmed[1:]),
                    ]
                    markdown = "\n".join(lines)
                    add(
                        TableNode,
                        text=markdown,
                        heading_path=path,
                        anchor=anchor,
                        markdown=markdown,
                        row_count=len(trimmed),
                        column_count=len(header),
                    )
                else:
                    text = "\n".join(" — ".join(c for c in r if c) for r in trimmed)
                    if text.strip():
                        add(ParagraphNode, text=text, heading_path=path, anchor=anchor)
        wb.close()
        return ParsedDocument(
            metadata=ParsedDocumentMetadata(
                source_id=source.source_id,
                filename=source.filename,
                parser_name=XLSX_PARSER_NAME,
                parser_version=openpyxl_version(),
                page_count=None,
                parse_warnings=tuple(warnings),
            ),
            nodes=tuple(nodes),
        )
