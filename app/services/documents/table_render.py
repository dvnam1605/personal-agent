"""Render a generated table to Markdown (UI preview) and .xlsx (download).

Layout of the workbook:

* sheet "Bảng": title, subtitle (date, number of source documents), header
  row (STT + columns + "Nguồn"), one row per table row; frozen header,
  wrapped text, widths per column; cells cleared by the source check are
  filled yellow with a comment saying why;
* sheet "Nguồn": the numbered source list the "Nguồn" column refers to.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.documents.render import _source_line

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_CLEARED_FILL = PatternFill("solid", fgColor="FFF2CC")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _refs(numbers: list[int]) -> str:
    return "".join(f"[{n}]" for n in numbers)


def _md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def render_table_markdown(
    title: str,
    columns: list[dict[str, Any]],
    rows: list[dict[str, Any]],
    citations: list[dict[str, Any]],
    *,
    created: date | None = None,
) -> str:
    created = created or date.today()
    names = ["STT", *[c["name"] for c in columns], "Nguồn"]
    parts = [
        f"# {title}",
        f"_Bảng tổng hợp — ngày {created:%d/%m/%Y} — {len(rows)} dòng_",
        "",
        "| " + " | ".join(names) + " |",
        "|" + "---|" * len(names),
    ]
    for i, row in enumerate(rows, 1):
        cells = [
            _md_cell(c.get("value") or ("(chưa xác minh)" if c.get("status") == "cleared" else ""))
            for c in row["cells"]
        ]
        parts.append("| " + " | ".join([str(i), *cells, _refs(row.get("citations", []))]) + " |")
    if not rows:
        parts.append(
            "| " + " | ".join(["", "Không trích được dữ liệu nào", *[""] * len(columns)]) + " |"
        )
    if citations:
        parts.extend(["", "## Nguồn", ""])
        parts.extend(f"[{c['number']}] {_source_line(c)}" for c in citations)
    return "\n".join(parts).strip() + "\n"


def _width(column: dict[str, Any]) -> int:
    source = column.get("source")
    if source in {"doc_number", "doc_date", "doc_type"}:
        return 16
    if source in {"signer", "issuer"}:
        return 24
    return 40


def render_table_xlsx(
    path: Path,
    title: str,
    columns: list[dict[str, Any]],
    rows: list[dict[str, Any]],
    citations: list[dict[str, Any]],
    *,
    request: str = "",
    created: date | None = None,
) -> Path:
    created = created or date.today()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Bảng"
    n_cols = len(columns) + 2  # STT + columns + Nguồn

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    subtitle = f"Lập ngày {created:%d/%m/%Y} từ {len({r['document_id'] for r in rows})} văn bản"
    if request:
        subtitle += f" — yêu cầu: {request}"
    ws.cell(row=2, column=1, value=subtitle).font = Font(italic=True, color="595959")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)

    header_row = 4
    headers = ["STT", *[c["name"] for c in columns], "Nguồn"]
    for j, name in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=j, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _BORDER

    for i, row in enumerate(rows, 1):
        r = header_row + i
        stt = ws.cell(row=r, column=1, value=i)
        stt.alignment = Alignment(horizontal="center", vertical="top")
        stt.border = _BORDER
        for j, data in enumerate(row["cells"], 2):
            cell = ws.cell(row=r, column=j, value=data.get("value") or None)
            cell.alignment = _WRAP_TOP
            cell.border = _BORDER
            if data.get("status") == "cleared":
                cell.fill = _CLEARED_FILL
                cell.comment = Comment(
                    "Chưa xác minh được với văn bản nguồn"
                    + (f": {data['note']}" if data.get("note") else ""),
                    "Trợ lý",
                )
        src = ws.cell(row=r, column=n_cols, value=_refs(row.get("citations", [])))
        src.alignment = _WRAP_TOP
        src.border = _BORDER

    ws.column_dimensions["A"].width = 6
    for j, column in enumerate(columns, 2):
        ws.column_dimensions[get_column_letter(j)].width = _width(column)
    ws.column_dimensions[get_column_letter(n_cols)].width = 12
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + len(rows)}"

    if any(c.get("status") == "cleared" for r in rows for c in r["cells"]):
        note_row = header_row + len(rows) + 2
        note = ws.cell(
            row=note_row,
            column=1,
            value="Ô tô vàng: dữ liệu chưa xác minh được với văn bản nguồn nên để trống, "
            "cần kiểm tra và điền tay.",
        )
        note.font = Font(italic=True, color="7F6000")
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=n_cols)

    src_ws = wb.create_sheet("Nguồn")
    for j, name in enumerate(("Số", "Văn bản / vị trí"), 1):
        cell = src_ws.cell(row=1, column=j, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
    for i, citation in enumerate(citations, 2):
        src_ws.cell(row=i, column=1, value=f"[{citation['number']}]")
        src_ws.cell(row=i, column=2, value=_source_line(citation)).alignment = _WRAP_TOP
    src_ws.column_dimensions["A"].width = 8
    src_ws.column_dimensions["B"].width = 90

    wb.save(path)
    return path
