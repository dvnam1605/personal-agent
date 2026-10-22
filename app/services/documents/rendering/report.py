"""Render a written document to Markdown (UI) and .docx (download).

The .docx uses python-docx (installed with docling). Administrative kinds
get the national motto header block; equation lines ("$$…$$ (3.5)") become
native Word equations; every document ends with a numbered "Tài liệu tham
khảo" list matching the inline [n] markers.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

import latex2mathml.converter
import mathml2omml
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Pt

from app.services.documents.templates import DocTemplate
from app.services.documents.writing.formulas import EQUATION_LINE

_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_TEXT_ARG = re.compile(r"\\text\{[^{}]*\}")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
# Underscores inside identifiers (insufficient_columns, MAX_EDIT_TIME) are not italics.
_ITALIC = re.compile(r"(?<![\w*])_(?=\S)(.+?)(?<=\S)_(?![\w*])")


def source_line(c: dict[str, Any]) -> str:
    bits = [str(c.get("title") or c.get("document_id"))]
    if c.get("section_title"):
        bits.append(f"mục {c['section_title']}")
    if c.get("page_start"):
        bits.append(f"tr. {c['page_start']}")
    return ", ".join(bits)


def render_markdown(
    title: str,
    sections: list[dict[str, Any]],
    citations: list[dict[str, Any]],
    *,
    template: DocTemplate,
    created: date | None = None,
) -> str:
    created = created or date.today()
    parts = [f"# {title}", f"_{template.label} — ngày {created:%d/%m/%Y}_", ""]
    for s in sections:
        if s.get("status") == "missing":
            continue
        parts.extend([f"## {s['heading']}", "", s["text"].strip(), ""])
    if citations:
        parts.extend(["## Tài liệu tham khảo", ""])
        parts.extend(f"[{c['number']}] {source_line(c)}" for c in citations)
    return "\n".join(parts).strip() + "\n"


def _add_runs(paragraph: Any, text: str) -> None:
    """Minimal inline Markdown: **bold** and _italic_."""
    pos = 0
    tokens: list[tuple[int, int, str, str]] = []
    for m in _BOLD.finditer(text):
        tokens.append((m.start(), m.end(), "b", m.group(1)))
    for m in _ITALIC.finditer(text):
        if not any(s <= m.start() < e for s, e, _, _ in tokens):
            tokens.append((m.start(), m.end(), "i", m.group(1)))
    for start, end, kind, inner in sorted(tokens):
        if start < pos:
            continue
        if start > pos:
            paragraph.add_run(text[pos:start])
        run = paragraph.add_run(inner)
        run.bold = kind == "b"
        run.italic = kind == "i"
        pos = end
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _m(tag: str) -> str:
    return f"{{{_MATH_NS}}}{tag}"


def _hide_empty_slots(math: Any) -> None:
    """Word draws a placeholder box for a missing limit or root degree unless hidden."""
    for nary in math.iter(_m("nary")):
        props = nary.find(_m("naryPr"))
        if props is None:
            props = nary.makeelement(_m("naryPr"), {})
            nary.insert(0, props)
        for slot, flag in (("sub", "subHide"), ("sup", "supHide")):
            part = nary.find(_m(slot))
            if part is not None and len(part) == 0:
                props.append(props.makeelement(_m(flag), {_m("val"): "1"}))
    for rad in math.iter(_m("rad")):
        if rad.find(_m("deg")) is None:
            props = rad.makeelement(_m("radPr"), {})
            props.append(props.makeelement(_m("degHide"), {_m("val"): "1"}))
            rad.insert(0, props)
            rad.insert(1, rad.makeelement(_m("deg"), {}))


def _omml(latex: str) -> Any | None:
    """LaTeX as a Word equation element (LaTeX → MathML → OMML), or None."""
    # latex2mathml keeps the backslash of "\_" inside \text{…}.
    latex = _TEXT_ARG.sub(lambda m: m.group(0).replace("\\_", "_"), latex)
    try:
        omml = mathml2omml.convert(latex2mathml.converter.convert(latex))
    except Exception:  # noqa: BLE001 - fall back to the LaTeX source as text
        return None
    math = parse_xml(omml.replace("<m:oMath>", f'<m:oMath xmlns:m="{_MATH_NS}">', 1))
    _hide_empty_slots(math)
    return math


def _add_equation(doc: Any, latex: str, number: str | None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    element = _omml(latex)
    if element is None:
        paragraph.add_run(latex)
    else:
        paragraph._p.append(element)
    if number:
        paragraph.add_run(f"    ({number})")


def _add_block(doc: Any, block: str) -> None:
    for raw in block.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.lstrip()
        if equation := EQUATION_LINE.match(stripped):
            _add_equation(doc, equation.group("latex").strip(), equation.group("num"))
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith(("- ", "* ", "• ")):
            _add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+[.)]\s", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s", "", stripped))
        elif stripped.startswith("|"):
            doc.add_paragraph(stripped)  # tables stay readable as text
        else:
            _add_runs(doc.add_paragraph(), stripped)


def render_docx(
    path: Path,
    title: str,
    sections: list[dict[str, Any]],
    citations: list[dict[str, Any]],
    *,
    template: DocTemplate,
    issuer: str | None = None,
    created: date | None = None,
) -> Path:
    created = created or date.today()
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)

    if template.formal_header:
        table = doc.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells
        left.text = (issuer or "CƠ QUAN, ĐƠN VỊ").upper()
        right.text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc"
        for cell in (left, right):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = True
        place = doc.add_paragraph(f"…, ngày {created.day} tháng {created.month} năm {created.year}")
        place.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    heading = doc.add_heading(template.label.upper() if template.formal_header else title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if template.formal_header:
        sub = doc.add_paragraph(title)
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in sub.runs:
            run.bold = True
    else:
        meta = doc.add_paragraph(f"{template.label} — ngày {created:%d/%m/%Y}")
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for s in sections:
        if s.get("status") == "missing":
            continue
        doc.add_heading(s["heading"], level=1)
        _add_block(doc, s["text"])

    if citations:
        doc.add_heading("Tài liệu tham khảo", level=1)
        for c in citations:
            doc.add_paragraph(f"[{c['number']}] {source_line(c)}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
