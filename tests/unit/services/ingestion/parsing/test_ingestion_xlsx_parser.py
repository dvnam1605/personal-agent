"""Excel workbooks: sheets → titles, label rows and tables with their header."""

from __future__ import annotations

import io

from openpyxl import Workbook

from app.domain.models.ingestion.documents import DetectedDocumentType, SourceDocument
from app.domain.models.ingestion.parsed_document import HeadingNode, ParagraphNode, TableNode
from app.services.ingestion.parsing.base import parse_source, parser_identity
from app.services.ingestion.type_detection import detect_document_type


def _workbook() -> bytes:
    wb = Workbook()
    log = wb.active
    assert log is not None
    log.title = "Quầy line"
    log.append(["ĐÁNH GIÁ KẾT QUẢ NHẬN DIỆN — QUẦY LINE"])
    log.append(["Từ 09/10/2026 đến 21/10/2026"])
    log.append([])
    log.append(["SL", "Thời Gian", "Độ chính xác", "Camera"])
    log.append([1, "2026-10-09 10:32:07", 1, "Aeon - Quầy line"])
    log.append([2, "2026-10-09 10:52:07", 0.75, "Aeon | Quầy line"])
    summary = wb.create_sheet("Kết quả chung")
    summary.append(["Quầy line"])
    summary.append(["Chỉ số", "Tử số", "Mẫu số", "Kết quả"])
    summary.append(["Tỷ lệ nhận diện chính xác", 15315, 15360, 15315 / 15360])
    wb.create_sheet("Trống")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _source(name: str = "Danh_gia.xlsx") -> SourceDocument:
    return SourceDocument(
        source_id="x1", filename=name, source_type="upload", checksum="0" * 64, size_bytes=1
    )


def test_xlsx_is_detected_from_zip_and_extension() -> None:
    data = _workbook()
    result = detect_document_type("Danh_gia.xlsx", None, data[:64])
    assert result.document_type is DetectedDocumentType.XLSX
    # A bare zip with no name/MIME hint stays unsupported.
    assert detect_document_type("blob", None, data[:64]).document_type is None


async def test_sheets_become_titled_tables() -> None:
    result = await parse_source(_source(), _workbook())
    assert result.document is not None
    nodes = result.document.nodes
    headings = [n.text for n in nodes if isinstance(n, HeadingNode)]
    assert headings == ["Quầy line", "Kết quả chung"]  # empty sheet skipped
    title = next(n for n in nodes if isinstance(n, ParagraphNode))
    assert "ĐÁNH GIÁ KẾT QUẢ NHẬN DIỆN" in title.text and "09/10/2026" in title.text
    tables = [n for n in nodes if isinstance(n, TableNode)]
    assert tables[0].markdown.splitlines()[0] == "| SL | Thời Gian | Độ chính xác | Camera |"
    assert "| 2 | 2026-10-09 10:52:07 | 0.75 | Aeon \\| Quầy line |" in tables[0].markdown
    assert tables[0].heading_path == ("Quầy line",)
    # "Quầy line" above the metrics is a label; the table starts at "Chỉ số".
    assert tables[1].markdown.startswith("| Chỉ số | Tử số | Mẫu số | Kết quả |")
    assert "| Tỷ lệ nhận diện chính xác | 15315 | 15360 | 0.99707 |" in tables[1].markdown
    assert parser_identity(_source(), _workbook())[0] == "openpyxl"
