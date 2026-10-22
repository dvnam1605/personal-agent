"""Source formulas are shown in the sections that present them."""

from __future__ import annotations

from app.services.documents.writing.formulas import collect_formulas, insert_formulas

TK = (
    "Với một thông số s, tập giá trị hợp lệ được xác định theo Công thức (3.5), còn khoảng "
    "giá trị tương ứng được biểu diễn trong Công thức (3.6):\n\n"
    "Công thức: Ps = {p ∈Pcategory | vs(p) có mặt, vs(p) ∈R} (3.5)\n\n"
    "Công thức: range(s) = [minp∈Ps vs(p), maxp∈Ps vs(p)], nếu |Ps| > 0 (3.6)"
)
RRF = (
    "Hai danh sách ứng viên được hợp nhất thứ hạng bằng Reciprocal Rank Fusion:\n\n"
    "Công thức: SRRF(d) = X i:d∈Li 1 k + ranki(d) (3.3)"
)
FORMULAS = collect_formulas(
    [
        ("3.4.2 Kỹ thuật TK-Table và Self-Augmented Prompting", TK),
        ("3.2 Hệ thống tìm kiếm lai", RRF),
    ]
)


def test_formulas_are_collected_with_their_section_and_context() -> None:
    assert FORMULAS["3.5"].body == "Ps = {p ∈Pcategory | vs(p) có mặt, vs(p) ∈R}"
    assert FORMULAS["3.6"].section.startswith("3.4.2")
    assert "Reciprocal Rank Fusion" in FORMULAS["3.3"].context


def test_mentioned_formulas_are_inserted_after_the_sentence() -> None:
    text = (
        "Tập giá trị hợp lệ được xác định theo Công thức (3.5), và khoảng giá trị được biểu "
        "diễn qua Công thức (3.6) [4]. Trong các biểu thức này, P_category là tập sản phẩm [4]."
    )
    out = insert_formulas(text, FORMULAS)
    assert out == (
        "Tập giá trị hợp lệ được xác định theo Công thức (3.5), và khoảng giá trị được biểu "
        "diễn qua Công thức (3.6) [4].\n\n"
        "Công thức (3.5): Ps = {p ∈Pcategory | vs(p) có mặt, vs(p) ∈R}\n\n"
        "Công thức (3.6): range(s) = [minp∈Ps vs(p), maxp∈Ps vs(p)], nếu |Ps| > 0\n\n"
        "Trong các biểu thức này, P_category là tập sản phẩm [4]."
    )


def test_formulas_of_the_planned_source_section_are_added_without_a_mention() -> None:
    text = (
        "Hệ thống dùng PostgreSQL và Qdrant [1]. Hai danh sách ứng viên được hợp nhất thứ hạng "
        "bằng Reciprocal Rank Fusion để ra danh sách cuối [2]. Bộ nhớ đệm lưu ngữ cảnh [3]."
    )
    out = insert_formulas(
        text + "\n\nBộ nhớ đệm Redis lưu ngữ cảnh hội thoại [3].",
        FORMULAS,
        sources=["3.2 Hệ thống tìm kiếm lai"],
    )
    assert (
        "Bộ nhớ đệm lưu ngữ cảnh [3].\n\n"
        "Hai danh sách ứng viên được hợp nhất thứ hạng bằng Reciprocal Rank Fusion:\n\n"
        "Công thức (3.3): SRRF(d) = X i:d∈Li 1 k + ranki(d)\n\n"
        "Bộ nhớ đệm Redis"
    ) in out
    assert "(3.5)" not in out  # not this section's source


def test_a_colon_does_not_end_the_mentioning_sentence() -> None:
    out = insert_formulas(
        "Điểm được tính theo Công thức (3.3) cho hai dạng: tên và mô tả [1]. Tiếp theo [2].",
        FORMULAS,
    )
    assert out.startswith(
        "Điểm được tính theo Công thức (3.3) cho hai dạng: tên và mô tả [1].\n\nCông thức (3.3)"
    )


def test_formula_already_written_or_unknown_is_left_alone() -> None:
    written = "Theo công thức (3.5): Ps = {p ∈Pcategory | vs(p) có mặt, vs(p) ∈R} [1]."
    assert insert_formulas(written, FORMULAS) == written
    assert insert_formulas("Theo Công thức (9.9) [1].", FORMULAS) == "Theo Công thức (9.9) [1]."


def test_each_formula_is_shown_once_in_the_document() -> None:
    shown: set[str] = set()
    a = insert_formulas("Công thức (3.5) dùng ở đây [1].", FORMULAS, already_shown=shown)
    b = insert_formulas("Nhắc lại công thức (3.5) [2].", FORMULAS, already_shown=shown)
    assert "Ps =" in a and "Ps =" not in b


def test_parent_section_number_covers_its_subsections() -> None:
    out = insert_formulas(
        "Kỹ thuật Self-Augmented Prompting bổ sung khoảng giá trị của thông số [1].",
        FORMULAS,
        sources=["3.4 Các kỹ thuật từ nghiên cứu khoa học"],
    )
    assert "Công thức (3.5)" in out and "Công thức (3.6)" in out and "(3.3)" not in out


def test_formula_without_a_clear_spot_goes_to_the_end_of_the_section() -> None:
    out = insert_formulas(
        "Hệ thống dùng PostgreSQL và Qdrant [1]. Bộ nhớ đệm lưu ngữ cảnh [2].",
        FORMULAS,
        sources=["3.2 Hệ thống tìm kiếm lai"],
    )
    assert out == (
        "Hệ thống dùng PostgreSQL và Qdrant [1]. Bộ nhớ đệm lưu ngữ cảnh [2].\n\n"
        "Hai danh sách ứng viên được hợp nhất thứ hạng bằng Reciprocal Rank Fusion:\n\n"
        "Công thức (3.3): SRRF(d) = X i:d∈Li 1 k + ranki(d)"
    )


def test_formulas_sharing_a_lead_are_introduced_once() -> None:
    out = insert_formulas(
        "Hệ thống dùng PostgreSQL [1].",
        FORMULAS,
        sources=["3.4 Các kỹ thuật từ nghiên cứu khoa học"],
    )
    assert out.count("Với một thông số s, tập giá trị hợp lệ") == 1
    assert (
        out.index("Với một thông số s")
        < out.index("\n\nCông thức (3.5):")
        < out.index("\n\nCông thức (3.6):")
    )


def test_formula_hint_lists_only_the_section_formulas() -> None:
    from app.services.documents.writing.formulas import formula_hint

    hint = formula_hint(FORMULAS, ["3.4.2 Kỹ thuật TK-Table và Self-Augmented Prompting"])
    assert "(3.5)" in hint and "(3.6)" in hint and "(3.3)" not in hint


COSINE = r"\cos(a, b) = \frac{a \cdot b}{\|a\| \|b\|}"
LATEX_CHUNK = (
    "Mức độ tương đồng được đo bằng cosine similarity theo Công thức (3.2):\n\n"
    f"Công thức (3.2): $${COSINE}$$"
)


def test_latex_formulas_are_collected_and_shown_as_equation_lines() -> None:
    formulas = collect_formulas([("3.2 Hệ thống tìm kiếm lai", LATEX_CHUNK)])
    assert formulas["3.2"].latex == COSINE
    assert "cosine similarity" in formulas["3.2"].context
    out = insert_formulas("Độ tương đồng tính theo Công thức (3.2) [1].", formulas)
    assert out.endswith(f"$${COSINE}$$ (3.2)")


def test_equation_lines_become_word_equations(tmp_path) -> None:
    from docx import Document

    from app.services.documents.rendering.report import render_docx
    from app.services.documents.templates import TEMPLATE_BY_KEY

    latex = r"\text{cos\_sim} = \sum_{i:d \in L_i} \frac{1}{\sqrt{k+r}}"
    text = f"Theo đồ án [1].\n\n$${latex}$$ (3.3)"
    path = render_docx(
        tmp_path / "f.docx",
        "Báo cáo",
        [{"heading": "I. Kỹ thuật", "text": text}],
        [],
        template=TEMPLATE_BY_KEY["summary_report"],
    )
    xml = Document(str(path)).element.xml
    assert "<m:oMath" in xml and "<m:f>" in xml
    assert "(3.3)" in xml and "$$" not in xml
    # No placeholder boxes for the missing upper limit and root degree.
    assert '<m:supHide m:val="1"/>' in xml and '<m:degHide m:val="1"/>' in xml
    assert "cos_sim" in xml and "cos\\_sim" not in xml


def test_identifiers_with_underscores_are_not_italicised(tmp_path) -> None:
    from docx import Document

    from app.services.documents.rendering.report import render_docx
    from app.services.documents.templates import TEMPLATE_BY_KEY

    text = "Lỗi insufficient_columns, giới hạn STRUCTURED_QA_MAX_EDIT_TIME = 3 và _nhấn mạnh_ [1]."
    path = render_docx(
        tmp_path / "u.docx",
        "Báo cáo",
        [{"heading": "I. Kỹ thuật", "text": text}],
        [],
        template=TEMPLATE_BY_KEY["summary_report"],
    )
    paragraph = next(p for p in Document(str(path)).paragraphs if "insufficient" in p.text)
    assert "insufficient_columns" in paragraph.text
    assert "STRUCTURED_QA_MAX_EDIT_TIME" in paragraph.text
    assert [r.text for r in paragraph.runs if r.italic] == ["nhấn mạnh"]


def test_inline_math_in_a_sentence_becomes_a_word_equation(tmp_path) -> None:
    from docx import Document

    from app.services.documents.rendering.report import render_docx
    from app.services.documents.templates import TEMPLATE_BY_KEY

    text = r"Trong đó, tập $P_s$ gồm các sản phẩm thuộc $\mathbb{R}$ [1]."
    path = render_docx(
        tmp_path / "m.docx",
        "Báo cáo",
        [{"heading": "I. Kỹ thuật", "text": text}],
        [],
        template=TEMPLATE_BY_KEY["summary_report"],
    )
    paragraph = next(p for p in Document(str(path)).paragraphs if "Trong đó" in p.text)
    xml = paragraph._p.xml
    assert xml.count("<m:oMath") == 2 and "$" not in paragraph.text
    assert paragraph.text.startswith("Trong đó, tập ") and paragraph.text.endswith("[1].")


def test_an_image_appears_once_in_the_whole_report() -> None:
    from app.services.documents.service import _unique_images

    picture = "![Hình 4.15: Giao diện quản trị](image:" + "e" * 24 + ")"
    sections = [
        {"heading": "I", "text": f"Tổng quan [1].\n{picture}"},
        {"heading": "II", "text": f"Chi tiết [2].\n{picture}\n\nKết luận [2]."},
    ]
    out = _unique_images(sections)
    assert picture in out[0]["text"]
    assert out[1]["text"] == "Chi tiết [2].\n\nKết luận [2]."


def test_a_section_is_offered_the_figures_of_its_source_headings() -> None:
    from app.core.images import image_ref
    from app.services.documents.writing.figures import collect_figures, figure_hint

    arch = image_ref("Hình 4.2: Sơ đồ kiến trúc — mô tả tự động: ba tầng", "1" * 24)
    ui = image_ref("Hình 4.12: Giao diện trang chủ", "2" * 24)
    figures = collect_figures(
        [("4.1 Kiến trúc tổng thể", f"Kiến trúc.\n\n{arch}"), ("4.4 Giao diện", f"UI.\n\n{ui}")]
    )
    hint, refs = figure_hint(figures, ["4.1 Kiến trúc tổng thể"])
    assert hint == "- [Hình 4.2] Sơ đồ kiến trúc — mô tả tự động: ba tầng"
    assert refs == arch


async def test_markers_resolve_against_the_offered_figures() -> None:
    import uuid

    from app.core.images import image_ref
    from app.domain.models.retrieval import Evidence, EvidenceBundle
    from app.services.retrieval.synthesis import PromptAnswerSynthesizer

    bundle = EvidenceBundle(
        items=[
            Evidence(
                evidence_id=uuid.uuid4().hex,
                kind="CHILD",
                content_raw="Hệ thống gồm ba tầng.",
                token_estimate=5,
                primary_chunk_id="c1",
                document_id="doc-1",
            )
        ],
        total_tokens=5,
        documents_used=["doc-1"],
        parent_ids_used=[],
        retrieval_trace_id=uuid.uuid4().hex,
    )
    arch = image_ref("Hình 4.2: Sơ đồ kiến trúc", "1" * 24)

    async def gen(system: str, user: str) -> str:
        return "Hệ thống gồm ba tầng [E1].\n[Hình 4.2]"

    result = await PromptAnswerSynthesizer(generate=gen).synthesize(
        "Viết mục kiến trúc", bundle, extra_images=arch
    )
    assert result.answer == f"Hệ thống gồm ba tầng [1].\n{arch}"


def test_inline_code_loses_its_backticks(tmp_path) -> None:
    from docx import Document

    from app.services.documents.rendering.report import render_docx
    from app.services.documents.templates import TEMPLATE_BY_KEY

    path = render_docx(
        tmp_path / "c.docx",
        "Báo cáo",
        [{"heading": "I", "text": "Thuộc tính `sql_client` truy vấn PostgreSQL [1]."}],
        [],
        template=TEMPLATE_BY_KEY["summary_report"],
    )
    paragraph = next(p for p in Document(str(path)).paragraphs if "Thuộc tính" in p.text)
    assert paragraph.text == "Thuộc tính sql_client truy vấn PostgreSQL [1]."
    assert [r.font.name for r in paragraph.runs if r.text == "sql_client"] == ["Consolas"]
