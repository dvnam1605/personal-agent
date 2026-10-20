"""Document generation: templates, outline, grounded sections, rendering, lifecycle."""

from __future__ import annotations

import uuid
from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxDocument
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.domain.models.retrieval.citation import Citation
from app.domain.models.retrieval.metadata_filter import (
    KnowledgeIntent,
    KnowledgePlan,
    MetadataFilterSpec,
)
from app.domain.models.retrieval.sufficiency import SufficiencyStatus
from app.infrastructure.db.base import Base
from app.infrastructure.db.models import Document, User
from app.services.documents.generation import (
    MISSING_SECTION_NOTE,
    DocumentOutline,
    DocumentPlanner,
    OutlineSection,
    SectionWriter,
    normalize_outline_sections,
)
from app.services.documents.render import render_docx, render_markdown
from app.services.documents.service import DocumentGenerationService, outline_message
from app.services.documents.templates import (
    TEMPLATE_BY_KEY,
    is_generation_request,
    select_template,
)
from app.services.retrieval.synthesis import SynthesisResult

# ------------------------------------------------------------- templates


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Viết báo cáo tuần cho nhóm AI", True),
        ("Soạn tờ trình xin kinh phí đào tạo", True),
        ("Hãy lập biên bản cuộc họp sáng nay", True),
        ("Tóm tắt giúp tôi đồ án", False),  # no generation verb → plain RAG
        ("Soạn email gửi anh Tuấn về báo cáo", False),  # e-mail drafting
        ("Báo cáo tuần trước nói gì?", False),
    ],
)
def test_is_generation_request(text: str, expected: bool) -> None:
    assert is_generation_request(text) is expected


def test_select_template() -> None:
    assert select_template("viết báo cáo tuần").key == "weekly_report"
    assert select_template("soạn tờ trình xin kinh phí").key == "proposal"
    assert select_template("lập biên bản họp giao ban").key == "meeting_minutes"
    assert select_template("viết tài liệu về quy trình").key == "summary_report"


# --------------------------------------------------------------- fakes


class FakeMetadata:
    def __init__(self, ids: list[str]) -> None:
        self.ids = ids

    async def plan(self, request: str, requester_id: str | None) -> KnowledgePlan:
        spec = MetadataFilterSpec(categories=("report",)) if self.ids else MetadataFilterSpec()
        return KnowledgePlan(intent=KnowledgeIntent.CONTENT, spec=spec)

    async def resolve_document_ids(self, spec: Any, requester_id: Any, *, cap: int):
        return self.ids[:cap], len(self.ids)


def _citation(doc: str, chunk: str) -> Citation:
    return Citation(evidence_id=uuid.uuid4().hex, document_id=doc, chunk_id=chunk, title=f"T-{doc}")


class FakePipeline:
    """Section i cites the chunks in ``plan[i]``; empty → no grounded text."""

    def __init__(self, plan: list[list[tuple[str, str]]]) -> None:
        self.plan = plan
        self.queries: list[Any] = []

    async def run(self, query: Any) -> Any:  # pragma: no cover - not used when ids exist
        raise AssertionError("metadata ids should be used")

    async def run_with_synthesis(
        self, query: Any, internal_only: bool = True, system_prompt: str | None = None
    ) -> SynthesisResult:
        idx = len(self.queries)
        self.queries.append(query)
        cites = [_citation(d, c) for d, c in self.plan[idx]]
        if not cites:
            return SynthesisResult(answer="⚠️ ...", citations=[], status=SufficiencyStatus.PARTIAL)
        answer = " ".join(f"Ý {n} [{n}]." for n in range(1, len(cites) + 1))
        evidence = [{"cited": True, "prompt_text": "nội dung"} for _ in cites]
        return SynthesisResult(answer=answer, citations=cites, evidence=evidence)


# -------------------------------------------------------------- planner


async def test_planner_uses_metadata_sources_and_template_outline() -> None:
    planner = DocumentPlanner(FakeMetadata(["doc-1", "doc-2"]), FakePipeline([]), use_llm=False)

    async def titles(ids: list[str]) -> list[str]:
        return [f"Tài liệu {i}" for i in ids]

    outline = await planner.plan("Viết báo cáo tuần cho nhóm AI", None, source_titles_for=titles)
    assert outline.template_key == "weekly_report"
    assert outline.source_document_ids == ["doc-1", "doc-2"]
    assert [s.heading for s in outline.sections] == [
        s.heading for s in TEMPLATE_BY_KEY["weekly_report"].sections
    ]
    assert outline.title.startswith("Báo cáo tuần")


def test_templates_have_no_generic_summary_heading() -> None:
    headings = [s.heading for t in TEMPLATE_BY_KEY.values() for s in t.sections]
    assert not any("tổng hợp" in h.lower() for h in headings)


def test_outline_normalizer_renames_generic_heading_and_renumbers() -> None:
    template = TEMPLATE_BY_KEY["summary_report"]
    proposed = [
        OutlineSection(heading="I. Đặc điểm, tình hình chung", purpose="a"),
        OutlineSection(heading="II. Nội dung tổng hợp", purpose="b"),
        OutlineSection(heading="2.1 Khen thưởng Liên hoan Phát thanh", purpose="c"),
        OutlineSection(heading="Chế độ nhân dịp 8/3", purpose="d"),
        OutlineSection(heading="V. Đánh giá chung", purpose="e"),
    ]
    out = normalize_outline_sections(template, proposed)
    assert [s.heading for s in out] == [
        "I. Đặc điểm, tình hình chung",
        "II. Kết quả thực hiện",
        "III. Khen thưởng Liên hoan Phát thanh",
        "IV. Chế độ nhân dịp 8/3",
        "V. Đánh giá chung",
    ]


# --------------------------------------------------------------- writer


async def test_writer_renumbers_citations_across_sections_and_marks_missing() -> None:
    outline = DocumentOutline(
        template_key="summary_report",
        title="Báo cáo thử",
        sections=[OutlineSection(heading=h, purpose="") for h in ("A", "B", "C")],
        source_document_ids=["d1", "d2"],
    )
    pipeline = FakePipeline(
        [
            [("d1", "c1"), ("d2", "c2")],  # A: [1] [2]
            [("d2", "c2"), ("d1", "c3")],  # B: c2 reused → [2], c3 new → [3]
            [],  # C: nothing grounded
        ]
    )
    written = await SectionWriter(pipeline).write(outline, "viết báo cáo", None)
    a, b, c = written.sections
    assert a.text == "Ý 1 [1]. Ý 2 [2]."
    assert b.text == "Ý 1 [2]. Ý 2 [3]."
    assert c.status == "missing" and c.text == MISSING_SECTION_NOTE
    assert [x["number"] for x in written.citations] == [1, 2, 3]
    # Every section query is scoped to the selected sources.
    assert all(q.document_ids == ["d1", "d2"] for q in pipeline.queries)


# --------------------------------------------------------------- render


def test_render_markdown_and_docx(tmp_path: Path) -> None:
    template = TEMPLATE_BY_KEY["proposal"]
    sections = [
        {"heading": "I. Sự cần thiết", "text": "Lý do **quan trọng** [1].\n- ý phụ [1]"},
        {"heading": "II. Nội dung đề xuất", "text": MISSING_SECTION_NOTE},
    ]
    citations = [{"number": 1, "title": "Quyết định 1", "page_start": 2}]
    md = render_markdown("Tờ trình thử", sections, citations, template=template)
    assert "## I. Sự cần thiết" in md and "[1] Quyết định 1, tr. 2" in md

    path = render_docx(tmp_path / "x.docx", "Tờ trình thử", sections, citations, template=template)
    doc = DocxDocument(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "I. Sự cần thiết" in texts and "Tài liệu tham khảo" in texts
    assert any("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in c.text for c in doc.tables[0].rows[0].cells)


# ------------------------------------------------------------ lifecycle


@pytest.fixture
async def session() -> AsyncIterator[AsyncSession]:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as s:
        s.add(User(id="u1", email="u1@x.local"))
        s.add_all(
            [
                Document(
                    id=f"doc-{i}",
                    logical_document_id=f"l{i}",
                    source_type="upload",
                    title=f"file-{i}.pdf",
                    status="active",
                    is_active=True,
                )
                for i in (1, 2)
            ]
        )
        await s.commit()
        yield s
    await engine.dispose()


async def test_outline_then_approve_writes_files(
    session: AsyncSession, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("GENERATED_DOCS_DIR", str(tmp_path))
    pipeline = FakePipeline([[("doc-1", "c1")], [("doc-2", "c2")], [], []])
    service = DocumentGenerationService(
        metadata_service=FakeMetadata(["doc-1", "doc-2"]), pipeline=pipeline, use_llm=False
    )
    row = await service.create_outline(
        session, user_id="u1", requester_id=None, request="Viết báo cáo tuần nhóm AI"
    )
    assert row.status == "outline"
    assert "file-1.pdf" in outline_message(row)

    done = await service.approve_and_write(session, user_id="u1", requester_id=None, doc_id=row.id)
    assert done.status == "done"
    assert Path(done.docx_path).exists()
    assert done.markdown and "## Tài liệu tham khảo" in done.markdown
    assert [s["status"] for s in done.sections] == ["ok", "ok", "missing", "missing"]

    # Another user cannot touch it.
    with pytest.raises(LookupError):
        await service.approve_and_write(session, user_id="u2", requester_id=None, doc_id=row.id)


async def test_many_candidates_are_ranked_by_relevance_not_recency() -> None:
    from types import SimpleNamespace

    ids = [f"doc-{i}" for i in range(12)]

    class RankingPipeline(FakePipeline):
        async def run(self, query: Any) -> Any:
            assert query.document_ids == ids  # RAG restricted to the candidates
            items = [SimpleNamespace(document_id=d) for d in ("doc-9", "doc-3", "doc-9")]
            return SimpleNamespace(items=items)

    planner = DocumentPlanner(FakeMetadata(ids), RankingPipeline([]), use_llm=False)
    sources = await planner.select_sources("Viết báo cáo về các quyết định năm 2026", None)
    assert sources[:2] == ["doc-9", "doc-3"]
    assert len(sources) == 8
